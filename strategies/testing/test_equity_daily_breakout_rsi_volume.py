import datetime as dt
import importlib.util
import pathlib
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook


REPO_ROOT = pathlib.Path(__file__).resolve().parents[2]
MODULE_PATH = (
    REPO_ROOT
    / "strategies"
    / "scripts"
    / "EquityDailyBreakoutRsiVolume15min.py"
)
sys.path.insert(0, str(REPO_ROOT / "skills" / "fyers-trading" / "scripts"))
sys.path.insert(0, str(REPO_ROOT / "strategies" / "config"))
SPEC = importlib.util.spec_from_file_location(
    "equity_daily_breakout_rsi_volume", MODULE_PATH
)
assert SPEC and SPEC.loader
strategy = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = strategy
SPEC.loader.exec_module(strategy)


def intraday_candles(closes):
    latest_start = dt.datetime(2026, 1, 2, 9, 45, tzinfo=strategy.MARKET_TIMEZONE)
    start = int(latest_start.timestamp()) - ((len(closes) - 1) * strategy.CANDLE_SECONDS)
    return [
        strategy.Candle(
            epoch=start + index * strategy.CANDLE_SECONDS,
            open=close - 1,
            high=close + 1,
            low=close - 2,
            close=close,
            volume=100_000,
        )
        for index, close in enumerate(closes)
    ]


def daily_candles(previous_day_close=120.0, volume=20_000_000.0):
    latest_day = dt.datetime(2026, 1, 1, tzinfo=strategy.MARKET_TIMEZONE)
    candles = []
    for index in range(20):
        day = latest_day - dt.timedelta(days=19 - index)
        close = previous_day_close - (19 - index) * 0.1
        candles.append(
            strategy.Candle(
                epoch=int(day.timestamp()),
                open=close - 1,
                high=close + 1,
                low=close - 2,
                close=close,
                volume=volume,
            )
        )
    return candles


class DailyBreakoutStrategyTests(unittest.TestCase):
    def test_all_conditions_match(self):
        intraday = intraday_candles([100 + index for index in range(50)])
        daily = daily_candles(previous_day_close=145.0)
        # Make the latest candle clearly above the previous daily close.
        intraday[-1] = strategy.Candle(
            epoch=intraday[-1].epoch,
            open=146.0,
            high=151.0,
            low=145.0,
            close=150.0,
            volume=100_000,
        )
        intraday[-2] = strategy.Candle(
            epoch=intraday[-2].epoch,
            open=140.0,
            high=146.0,
            low=139.0,
            close=145.0,
            volume=100_000,
        )

        matched, details = strategy.evaluate_strategy(intraday, daily)

        self.assertTrue(matched, details)
        self.assertTrue(details["conditions"]["close_or_open_above_previous_day_close"])
        self.assertTrue(details["conditions"]["rsi_above_55"])
        self.assertTrue(details["conditions"]["average_daily_volume_above_10_million"])
        self.assertTrue(details["conditions"]["price_above_100"])
        self.assertEqual(details["previous_day_close"], 145.0)
        self.assertIn("rsi_crossed_above_55", details)

    def test_failed_conditions_are_reported(self):
        intraday = intraday_candles([100 + index for index in range(50)])
        daily = daily_candles(previous_day_close=200.0, volume=1_000_000.0)
        intraday[-1] = strategy.Candle(
            epoch=intraday[-1].epoch,
            open=110.0,
            high=111.0,
            low=109.0,
            close=110.0,
            volume=1,
        )

        matched, details = strategy.evaluate_strategy(intraday, daily)

        self.assertFalse(matched)
        self.assertIn(
            "close_or_open_above_previous_day_close",
            details["failed_conditions"],
        )
        self.assertIn(
            "average_daily_volume_above_10_million",
            details["failed_conditions"],
        )

    def test_todays_daily_candle_is_excluded(self):
        current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
        prior_epoch = int(dt.datetime(2026, 1, 1, tzinfo=current.tzinfo).timestamp())
        current_epoch = int(dt.datetime(2026, 1, 2, tzinfo=current.tzinfo).timestamp())
        response = {
            "candles": [
                [current_epoch, 101, 102, 100, 101, 20_000_000],
                [prior_epoch, 100, 101, 99, 100, 20_000_000],
            ]
        }

        candles = strategy.completed_daily_candles(response, current)

        self.assertEqual([candle.epoch for candle in candles], [prior_epoch])

    def test_incomplete_15_minute_candle_is_excluded(self):
        current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
        complete_epoch = int(current.timestamp()) - strategy.CANDLE_SECONDS
        incomplete_epoch = complete_epoch + 1
        response = {
            "candles": [
                [incomplete_epoch, 100, 101, 99, 100, 1],
                [complete_epoch, 100, 101, 99, 100, 1],
            ]
        }

        candles = strategy.completed_intraday_candles(response, current)

        self.assertEqual([candle.epoch for candle in candles], [complete_epoch])


class OrderAndLedgerTests(unittest.TestCase):
    def test_dry_run_match_is_printed_and_saved_once(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append((order, dry_run, meta, retry_transient))
                return {"s": "dry_run", "code": 0, "message": "dry-run"}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            excel_path = root / "signals.xlsx"
            client = FakeClient()
            intraday = intraday_candles([100 + index for index in range(50)])
            daily = daily_candles(previous_day_close=145.0)
            intraday[-1] = strategy.Candle(
                epoch=intraday[-1].epoch,
                open=146.0,
                high=151.0,
                low=145.0,
                close=150.0,
                volume=100_000,
            )
            intraday[-2] = strategy.Candle(
                epoch=intraday[-2].epoch,
                open=140.0,
                high=146.0,
                low=139.0,
                close=145.0,
                volume=100_000,
            )

            with (
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message") as log_message,
            ):
                first = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", intraday, daily, False
                )
                second = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", intraday, daily, False
                )
            strategy.flush_excel(store, excel_path)

            self.assertTrue(first)
            self.assertFalse(second)
            self.assertEqual(len(client.orders), 1)
            self.assertEqual(client.orders[0][0]["qty"], 7)
            self.assertEqual(client.orders[0][0]["type"], strategy.ORDER_TYPE)
            self.assertEqual(client.orders[0][0]["side"], strategy.ORDER_SIDE)
            self.assertEqual(client.orders[0][0]["productType"], strategy.PRODUCT_TYPE)
            self.assertTrue(client.orders[0][1])
            self.assertFalse(client.orders[0][3])
            self.assertTrue(
                any(
                    str(call.args[0]).startswith("MATCH |")
                    for call in log_message.call_args_list
                )
            )

            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active.max_row, 2)
            self.assertEqual(workbook.active["E2"].value, "NSE:TEST-EQ")
            self.assertEqual(workbook.active["V2"].value, "DRY_RUN")
            workbook.close()

    def test_ambiguous_live_order_is_reconciled_without_resend(self):
        class FakeClient:
            def __init__(self, tag):
                self.tag = tag
                self.orders = []
                self.orderbook_calls = 0

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append(order)
                raise strategy.AmbiguousOrderError("connection lost")

            def orderbook(self):
                self.orderbook_calls += 1
                return {
                    "s": "ok",
                    "orderBook": [{
                        "id": "OID123",
                        "orderTag": self.tag,
                        "orderStatus": "TRADED",
                    }],
                }

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            client = FakeClient("placeholder")
            intraday = intraday_candles([100 + index for index in range(50)])
            daily = daily_candles(previous_day_close=145.0)
            intraday[-1] = strategy.Candle(
                epoch=intraday[-1].epoch,
                open=146.0,
                high=151.0,
                low=145.0,
                close=150.0,
                volume=100_000,
            )
            intraday[-2] = strategy.Candle(
                epoch=intraday[-2].epoch,
                open=140.0,
                high=146.0,
                low=139.0,
                close=145.0,
                volume=100_000,
            )
            current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
            signal_id = (
                f"{strategy.STRATEGY_NAME}:NSE:TEST-EQ:{intraday[-1].epoch}"
            )
            client.tag = strategy.order_tag_for_signal(signal_id)
            with (
                patch.object(strategy, "now_ist", return_value=current),
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message"),
            ):
                strategy.handle_match(
                    client,
                    store,
                    "NSE:TEST-EQ",
                    intraday,
                    daily,
                    True,
                )

            row = store.get_row(signal_id)
            self.assertEqual(row["order_status"], "RECONCILED")
            self.assertEqual(row["order_id"], "OID123")
            self.assertEqual(len(client.orders), 1)
            self.assertEqual(client.orderbook_calls, 1)

    def test_insufficient_history_does_not_create_signal(self):
        with tempfile.TemporaryDirectory() as directory:
            store = strategy.SignalStore(pathlib.Path(directory) / "signals.db")
            matched, details = strategy.evaluate_strategy([], [])

            self.assertFalse(matched)
            self.assertIn("reason", details)
            self.assertEqual(store.pending_excel_rows(), [])


if __name__ == "__main__":
    unittest.main()
