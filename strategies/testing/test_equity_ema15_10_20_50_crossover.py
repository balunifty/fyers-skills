import datetime as dt
import importlib.util
import pathlib
import sys
import tempfile
import threading
import unittest
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


REPO_ROOT = pathlib.Path(__file__).resolve().parents[2]
MODULE_PATH = (
    REPO_ROOT
    / "strategies"
    / "scripts"
    / "EquityEma15_10_20_50Crossover15min.py"
)
sys.path.insert(0, str(REPO_ROOT / "skills" / "fyers-trading" / "scripts"))
sys.path.insert(0, str(REPO_ROOT / "strategies" / "config"))
SPEC = importlib.util.spec_from_file_location(
    "equity_ema15_10_20_50_crossover", MODULE_PATH
)
assert SPEC and SPEC.loader
strategy = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = strategy
SPEC.loader.exec_module(strategy)


def candles_from_closes(closes, latest_start=None):
    if latest_start is None:
        latest_start = dt.datetime(2026, 1, 2, 9, 45, tzinfo=strategy.MARKET_TIMEZONE)
    start = int(
        latest_start.timestamp()
    ) - ((len(closes) - 1) * strategy.CANDLE_SECONDS)
    return [
        strategy.Candle(
            epoch=start + index * strategy.CANDLE_SECONDS,
            open=close,
            high=close + 1,
            low=close - 1,
            close=close,
            volume=1000,
        )
        for index, close in enumerate(closes)
    ]


class FreshCrossoverTests(unittest.TestCase):
    def test_fresh_crossover_matches_all_conditions(self):
        candles = candles_from_closes([100.0] * 50 + [120.0])
        matched, details = strategy.evaluate_fresh_crossover(candles)

        self.assertTrue(matched, details)
        self.assertTrue(all(details["conditions"].values()))
        self.assertEqual(details["close"], 120.0)

    def test_ema10_pullback_cross_matches_upward_only_strategy(self):
        candles = candles_from_closes(
            [100.0] * 50 + [99.0, 101.0]
        )

        matched, details = strategy.evaluate_ema10_pullback_cross(candles)

        self.assertTrue(matched, details)
        self.assertTrue(details["conditions"]["previous_close_below_ema10"])
        self.assertTrue(details["conditions"]["current_close_above_ema10"])
        self.assertTrue(details["conditions"]["not_crossed_from_above_to_below_ema10"])
        self.assertTrue(details["conditions"]["ema10_ema20_ema50_bullish_stack"])

    def test_ema10_pullback_cross_rejects_downward_cross(self):
        candles = candles_from_closes(
            [100.0] * 50 + [101.0, 99.0]
        )

        matched, details = strategy.evaluate_ema10_pullback_cross(candles)

        self.assertFalse(matched)
        self.assertFalse(
            details["conditions"]["not_crossed_from_above_to_below_ema10"]
        )

    def test_already_bullish_ema_stack_is_rejected(self):
        candles = candles_from_closes([100.0] * 50 + [110.0, 120.0])
        matched, details = strategy.evaluate_fresh_crossover(candles)

        self.assertFalse(matched)
        self.assertFalse(
            details["conditions"][
                "not_previously_ema10_ema20_ema50_bullish_stack"
            ]
        )

    def test_incomplete_candle_is_removed(self):
        current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
        complete_epoch = int(current.timestamp()) - strategy.CANDLE_SECONDS
        incomplete_epoch = int(current.timestamp()) - strategy.CANDLE_SECONDS + 1
        response = {
            "candles": [
                [incomplete_epoch, 100, 101, 99, 100, 10],
                [complete_epoch, 99, 101, 98, 100, 10],
            ]
        }

        candles = strategy.completed_candles_from_response(response, current)

        self.assertEqual([candle.epoch for candle in candles], [complete_epoch])

    def test_delayed_latest_candle_is_rejected(self):
        current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
        delayed = strategy.Candle(
            epoch=int(current.timestamp()) - (2 * strategy.CANDLE_SECONDS),
            open=100,
            high=101,
            low=99,
            close=100,
            volume=1,
        )

        self.assertFalse(strategy.candle_is_current(delayed, current))

    def test_polling_waits_until_next_candle_boundary(self):
        current = dt.datetime(2026, 1, 2, 9, 27, 30, tzinfo=strategy.MARKET_TIMEZONE)

        self.assertEqual(
            strategy.seconds_until_next_candle_close(current),
            150,
        )


class ExactlyOnceTests(unittest.TestCase):
    def make_signal(
        self,
        store,
        mode="DRY_RUN",
        candle_epoch=None,
        strategy_name=None,
    ):
        if candle_epoch is None:
            candle_epoch = int(
                dt.datetime(2026, 1, 2, 9, 45, tzinfo=strategy.MARKET_TIMEZONE).timestamp()
            )
        details = {
            "candle_epoch": candle_epoch,
            "open": 100.0,
            "high": 121.0,
            "low": 99.0,
            "close": 120.0,
            "previous_close": 100.0,
            "ema15": 102.5,
            "ema10": 103.64,
            "ema20": 101.90,
            "ema50": 100.78,
            "conditions": {"fresh": True},
        }
        return store.claim(
            "NSE:TEST-EQ",
            details["candle_epoch"],
            "2027-01-01T09:30:00+05:30",
            "2027-01-01T09:46:00+05:30",
            details,
            mode,
            7,
            strategy_name or strategy.STRATEGY_NAME,
        )

    def test_ledger_and_workbook_each_accept_one_signal(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store)
            duplicate_id = self.make_signal(store)

            self.assertIsNotNone(signal_id)
            self.assertIsNone(duplicate_id)
            store.finish_order(signal_id, "DRY_RUN", "DRY_RUN", "nothing sent")
            row = store.get_row(signal_id)
            excel_path = root / "signals.xlsx"

            self.assertTrue(strategy.append_signal_to_excel(excel_path, row))
            self.assertTrue(strategy.append_signal_to_excel(excel_path, row))
            store.mark_excel_synced(signal_id, row["excel_revision"])
            strategy.flush_pending_excel(store, excel_path)

            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active.max_row, 2)
            self.assertEqual(workbook.active["A2"].value, signal_id)
            workbook.close()
            self.assertEqual(store.pending_excel_rows(), [])

    def test_same_candle_places_and_logs_only_one_order(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append((order, dry_run, meta, retry_transient))
                return {"s": "dry_run", "code": 0, "message": "dry-run"}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            excel_path = root / "signals.xlsx"
            client = FakeClient()
            candles = candles_from_closes([100.0] * 50 + [120.0])

            with (
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message") as log_message,
            ):
                first = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", candles, False, excel_path
                )
                second = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", candles, False, excel_path
                )

            self.assertTrue(first)
            self.assertFalse(second)
            self.assertEqual(len(client.orders), 1)
            self.assertEqual(client.orders[0][0]["qty"], 7)
            self.assertEqual(client.orders[0][0]["type"], strategy.ORDER_TYPE)
            self.assertTrue(client.orders[0][0]["orderTag"].startswith("e15x"))
            self.assertTrue(client.orders[0][1])
            self.assertFalse(client.orders[0][3])
            match_logs = [
                call for call in log_message.call_args_list
                if call.args and str(call.args[0]).startswith("MATCH |")
            ]
            self.assertEqual(len(match_logs), 1)
            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active.max_row, 2)
            workbook.close()

    def test_two_strategy_matches_get_independent_orders_and_rows(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append((order, meta))
                return {"s": "dry_run", "code": 0, "message": "dry-run"}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            client = FakeClient()
            excel_path = root / "signals.xlsx"
            candles = candles_from_closes([100.0] * 50 + [99.0, 101.0])

            with (
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message"),
            ):
                first = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", candles, False, excel_path
                )
                second = strategy.handle_match(
                    client,
                    store,
                    "NSE:TEST-EQ",
                    candles,
                    False,
                    excel_path,
                    strategy_name=strategy.STRATEGY_NAME_PULLBACK,
                    evaluator=strategy.evaluate_ema10_pullback_cross,
                )

            self.assertTrue(first)
            self.assertTrue(second)
            self.assertEqual(len(client.orders), 2)
            self.assertNotEqual(
                client.orders[0][0]["orderTag"], client.orders[1][0]["orderTag"]
            )
            self.assertEqual(
                {order[1]["strategy"] for order in client.orders},
                {strategy.STRATEGY_NAME, strategy.STRATEGY_NAME_PULLBACK},
            )
            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active.max_row, 3)
            workbook.close()

    def test_ambiguous_live_transport_is_never_resent(self):
        class FakeClient:
            def __init__(self):
                self.orders = []
                self.orderbook_calls = 0

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append(order)
                raise strategy.AmbiguousOrderError("connection lost after send")

            def orderbook(self):
                self.orderbook_calls += 1
                return {"s": "ok", "orderBook": []}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            excel_path = root / "signals.xlsx"
            client = FakeClient()
            candles = candles_from_closes([100.0] * 50 + [120.0])
            market_time = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)

            with (
                patch.object(strategy, "now_ist", return_value=market_time),
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message"),
            ):
                first = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", candles, True, excel_path
                )
                second = strategy.handle_match(
                    client, store, "NSE:TEST-EQ", candles, True, excel_path
                )

            self.assertTrue(first)
            self.assertFalse(second)
            self.assertEqual(len(client.orders), 1)
            self.assertGreaterEqual(client.orderbook_calls, 1)
            signal_id = (
                f"{strategy.STRATEGY_NAME}:NSE:TEST-EQ:{candles[-1].epoch}"
            )
            self.assertEqual(store.get_row(signal_id)["order_status"], "UNKNOWN")

    def test_matched_live_row_is_resumed_after_pre_send_crash(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append((order, dry_run, retry_transient))
                return {"s": "ok", "code": 1101, "id": "OID123", "message": "submitted"}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store, mode="LIVE")
            client = FakeClient()
            market_time = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)

            with (
                patch.object(strategy, "now_ist", return_value=market_time),
                patch.object(strategy, "log_message"),
            ):
                strategy.recover_unfinished_signals(
                    client, store, root / "signals.xlsx"
                )

            row = store.get_row(signal_id)
            self.assertEqual(row["order_status"], "PLACED")
            self.assertEqual(row["order_id"], "OID123")
            self.assertEqual(len(client.orders), 1)
            self.assertFalse(client.orders[0][2])

    def test_dry_run_invocation_cannot_send_persisted_live_signal(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, *args, **kwargs):
                self.orders.append((args, kwargs))
                return {"s": "ok", "code": 1101, "id": "OID123"}

            def orderbook(self):
                return {"s": "ok", "orderBook": []}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store, mode="LIVE")
            client = FakeClient()
            market_time = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)

            with (
                patch.object(strategy, "now_ist", return_value=market_time),
                patch.object(strategy, "log_message"),
            ):
                strategy.recover_unfinished_signals(
                    client,
                    store,
                    root / "signals.xlsx",
                    allow_live_send=False,
                )
            self.assertEqual(client.orders, [])
            self.assertEqual(store.get_row(signal_id)["order_status"], "MATCHED")

            with (
                patch.object(strategy, "now_ist", return_value=market_time),
                patch.object(strategy, "log_message"),
            ):
                strategy.recover_unfinished_signals(
                    client,
                    store,
                    root / "signals.xlsx",
                    allow_live_send=True,
                )
            self.assertEqual(len(client.orders), 1)
            self.assertEqual(store.get_row(signal_id)["order_status"], "PLACED")

    def test_previous_day_unfinished_signal_is_not_sent_on_recovery(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, *args, **kwargs):
                self.orders.append((args, kwargs))
                raise AssertionError("a stale signal must not be sent")

            def orderbook(self):
                return {"s": "ok", "orderBook": []}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(
                store, mode="LIVE", candle_epoch=1_800_000
            )
            client = FakeClient()
            with patch.object(strategy, "log_message"):
                strategy.recover_unfinished_signals(
                    client, store, root / "signals.xlsx"
                )

            row = store.get_row(signal_id)
            self.assertEqual(row["order_status"], "SKIPPED_STALE_RECOVERY")
            self.assertEqual(client.orders, [])

    def test_stale_same_day_candle_is_not_sent_on_recovery(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, *args, **kwargs):
                self.orders.append((args, kwargs))
                raise AssertionError("a stale same-day signal must not be sent")

            def orderbook(self):
                return {"s": "ok", "orderBook": []}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            old_epoch = int(
                dt.datetime(2026, 1, 2, 9, 30, tzinfo=strategy.MARKET_TIMEZONE).timestamp()
            )
            signal_id = self.make_signal(
                store, mode="LIVE", candle_epoch=old_epoch
            )
            client = FakeClient()
            current = dt.datetime(2026, 1, 2, 10, 0, tzinfo=strategy.MARKET_TIMEZONE)
            with (
                patch.object(strategy, "now_ist", return_value=current),
                patch.object(strategy, "log_message"),
            ):
                strategy.recover_unfinished_signals(
                    client, store, root / "signals.xlsx"
                )

            self.assertEqual(
                store.get_row(signal_id)["order_status"],
                "SKIPPED_STALE_RECOVERY",
            )
            self.assertEqual(client.orders, [])

    def test_submitting_row_without_tagged_order_becomes_reviewable_unknown(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, *args, **kwargs):
                self.orders.append((args, kwargs))
                raise AssertionError("an ambiguous row must never be resent")

            def orderbook(self):
                return {"s": "ok", "orderBook": []}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store, mode="LIVE")
            self.assertTrue(store.mark_submitting(signal_id))
            client = FakeClient()

            with patch.object(strategy, "log_message"):
                strategy.recover_unfinished_signals(
                    client, store, root / "signals.xlsx"
                )

            row = store.get_row(signal_id)
            self.assertEqual(row["order_status"], "UNKNOWN")
            self.assertEqual(row["order_id"], "RECONCILE_REQUIRED")
            self.assertEqual(client.orders, [])

    def test_submitting_row_is_reconciled_without_resending(self):
        class FakeClient:
            def __init__(self, tag):
                self.tag = tag
                self.orders = []

            def place_order(self, *args, **kwargs):
                self.orders.append((args, kwargs))
                raise AssertionError("an ambiguous row must never be resent")

            def orderbook(self):
                return {
                    "s": "ok",
                    "orderBook": [{
                        "id": "OID456",
                        "orderTag": self.tag,
                        "orderStatus": "TRADED",
                    }],
                }

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store, mode="LIVE")
            row = store.get_row(signal_id)
            self.assertTrue(store.mark_submitting(signal_id))
            client = FakeClient(row["order_tag"])

            with patch.object(strategy, "log_message"):
                strategy.recover_unfinished_signals(
                    client, store, root / "signals.xlsx"
                )

            row = store.get_row(signal_id)
            self.assertEqual(row["order_status"], "RECONCILED")
            self.assertEqual(row["order_id"], "OID456")
            self.assertEqual(client.orders, [])

    def test_excel_sync_does_not_clear_newer_sqlite_revision(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store)
            stale_row = store.get_row(signal_id)
            store.finish_order(signal_id, "PLACED", "OID789", "submitted")
            excel_path = root / "signals.xlsx"

            strategy.upsert_signal_to_excel(excel_path, stale_row)
            self.assertFalse(
                store.mark_excel_synced(signal_id, stale_row["excel_revision"])
            )
            self.assertEqual(len(store.pending_excel_rows()), 1)
            with patch.object(strategy, "log_message"):
                strategy.flush_pending_excel(store, excel_path)

            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active["T2"].value, "PLACED")
            self.assertEqual(workbook.active["U2"].value, "OID789")
            workbook.close()
            self.assertEqual(store.pending_excel_rows(), [])

    def test_legacy_excel_header_is_migrated(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            excel_path = root / "legacy.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(strategy.LEGACY_EXCEL_HEADERS)
            workbook.save(excel_path)
            workbook.close()

            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store)
            store.finish_order(signal_id, "DRY_RUN", "DRY_RUN", "nothing sent")
            strategy.upsert_signal_to_excel(excel_path, store.get_row(signal_id))

            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active["V1"].value, "OrderTag")
            self.assertEqual(workbook.active["X1"].value, "Details")
            self.assertTrue(str(workbook.active["V2"].value).startswith("e15x"))
            workbook.close()

    def test_concurrent_excel_flush_keeps_one_row_per_signal(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            database_path = root / "signals.db"
            excel_path = root / "signals.xlsx"
            store = strategy.SignalStore(database_path)
            first_id = self.make_signal(store, candle_epoch=1_800_000)
            second_id = self.make_signal(store, candle_epoch=1_800_001)
            store.finish_order(first_id, "DRY_RUN", "DRY_RUN", "first")
            store.finish_order(second_id, "DRY_RUN", "DRY_RUN", "second")
            stores = [
                strategy.SignalStore(database_path),
                strategy.SignalStore(database_path),
            ]
            barrier = threading.Barrier(2)

            def flush(worker_store):
                barrier.wait()
                return strategy.flush_pending_excel(worker_store, excel_path)

            with patch.object(strategy, "log_message"):
                with ThreadPoolExecutor(max_workers=2) as executor:
                    results = list(executor.map(flush, stores))

            self.assertEqual(results, [True, True])
            workbook = load_workbook(excel_path)
            self.assertEqual(workbook.active.max_row, 3)
            ids = {workbook.active.cell(row, 1).value for row in (2, 3)}
            self.assertEqual(ids, {first_id, second_id})
            workbook.close()

    def test_open_excel_defers_row_without_stopping_audit_state(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            signal_id = self.make_signal(store)
            store.finish_order(signal_id, "DRY_RUN", "DRY_RUN", "nothing sent")
            excel_path = root / "signals.xlsx"

            with (
                patch.object(strategy.os, "replace", side_effect=PermissionError("open")),
                patch.object(strategy, "log_message") as log_message,
            ):
                synced = strategy.flush_pending_excel(store, excel_path)

            self.assertFalse(synced)
            self.assertEqual(len(store.pending_excel_rows()), 1)
            self.assertTrue(log_message.called)
            with patch.object(strategy, "log_message"):
                self.assertTrue(strategy.flush_pending_excel(store, excel_path))
            self.assertEqual(store.pending_excel_rows(), [])

    def test_live_signal_after_entry_cutoff_is_logged_without_order(self):
        class FakeClient:
            def __init__(self):
                self.orders = []

            def place_order(self, order, dry_run, meta, retry_transient):
                self.orders.append(order)
                return {"s": "ok", "id": "SHOULD-NOT-HAPPEN"}

        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            store = strategy.SignalStore(root / "signals.db")
            client = FakeClient()
            candles = candles_from_closes(
                [100.0] * 50 + [120.0],
                dt.datetime(2026, 1, 2, 15, 0, tzinfo=strategy.MARKET_TIMEZONE),
            )
            late = dt.datetime(2026, 1, 2, 15, 25, tzinfo=strategy.MARKET_TIMEZONE)

            with (
                patch.object(strategy, "now_ist", return_value=late),
                patch.object(strategy, "get_entry_qty", return_value=7),
                patch.object(strategy, "log_message"),
            ):
                matched = strategy.handle_match(
                    client,
                    store,
                    "NSE:TEST-EQ",
                    candles,
                    True,
                    root / "signals.xlsx",
                )

            self.assertTrue(matched)
            self.assertEqual(client.orders, [])
            pending = store.pending_excel_rows()
            self.assertEqual(pending, [])
            workbook = load_workbook(root / "signals.xlsx")
            self.assertEqual(workbook.active["T2"].value, "SKIPPED_LATE_ENTRY")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
