"""Excel logger for strategy signals and order status."""

import datetime as dt
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

STRATEGIES_DIR = Path(__file__).resolve().parent.parent
EXCEL_PATH = STRATEGIES_DIR / "logs" / "Strategy_Trade_Log.xlsx"
HEADERS = ["DateTime", "Symbol", "StrategyName", "ScriptName", "SignalType", "OrderStatus", "OrderID", "Details"]
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _ensure_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Trade_Log"
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 15
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 50
    wb.save(path)
    return wb


def log_to_excel(strategy_name: str, script_name: str, symbol: str,
                 signal_type: str, order_status: str, order_id: str = "",
                 details: str = "") -> None:
    """Append a row to the Excel trade log.

    Args:
        strategy_name: e.g. "EMA_10_20_30_5MIN"
        script_name: e.g. "BuyCallOption102030EmaCrossover5min.py"
        symbol: stock or option symbol
        signal_type: e.g. "ENTRY", "EXIT", "SIGNAL_MATCH"
        order_status: e.g. "PLACED", "SKIPPED", "FAILED", "DRY_RUN"
        order_id: order ID if placed
        details: any extra info
    """
    EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = _ensure_workbook(EXCEL_PATH)
    ws = wb.active
    row = ws.max_row + 1
    timestamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    values = [timestamp, symbol, strategy_name, script_name, signal_type, order_status, order_id, details]
    for col_idx, value in enumerate(values, 1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.alignment = Alignment(horizontal="left")
    wb.save(EXCEL_PATH)
