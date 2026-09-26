#!/usr/bin/env python3
"""Buy NIFTY 50 equities on a fresh 15-minute EMA crossover setup.

Entry conditions are evaluated only on the latest completed 15-minute candle:
  1. Close crosses above EMA(15), and the current close is above EMA(10).
  2. EMA(10) crosses above EMA(20).
  3. EMA(20) crosses above EMA(50).
  4. Before the signal candle, EMA(10) > EMA(20) > EMA(50) must be false.

A second, independent strategy is also evaluated:
  1. The latest completed 15-minute candle moves from below EMA(10) to above
     EMA(10); a cross from above to below is rejected.
  2. The current close is above EMA(10).
  3. The current EMA stack is EMA(10) > EMA(20) > EMA(50).

Each strategy is independently deduplicated and may place its own order when
both match the same stock/candle. Every matched candle is claimed in SQLite
before an order attempt. A matched stock is printed and written to one
dedicated Excel row once, even if the
process is restarted or polls the same candle again. Live writes use a
per-signal orderTag, disable transient POST retries, and reconcile ambiguous
sends through the FYERS orderbook instead of resending them.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import importlib.util
import json
import os
import pathlib
import sqlite3
import sys
import tempfile
import time
from collections.abc import Callable
from contextlib import closing, contextmanager
from dataclasses import dataclass
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_DIR = pathlib.Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parents[1]
STRATEGIES_DIR = SCRIPT_DIR.parent

STOCKS_PATH = STRATEGIES_DIR / "data" / "Nifty50.txt"
DATABASE_PATH = (
    STRATEGIES_DIR
    / "databases"
    / "equity_ema15_10_20_50_crossover_15min.db"
)
LOG_PATH = STRATEGIES_DIR / "logs" / "EquityEma15_10_20_50Crossover15min.log"
EXCEL_PATH = (
    STRATEGIES_DIR / "logs" / "EquityEma15_10_20_50Crossover15min.xlsx"
)

STRATEGY_NAME = "EQUITY_EMA15_10_20_50_FRESH_CROSS_15MIN"
STRATEGY_NAME_PULLBACK = "EQUITY_EMA10_PULLBACK_UP_CROSS_15MIN"
SCRIPT_NAME = "EquityEma15_10_20_50Crossover15min.py"

MARKET_TIMEZONE = ZoneInfo("Asia/Kolkata")
MARKET_OPEN = dt.time(9, 15)
MARKET_CLOSE = dt.time(15, 30)
LAST_INTRADAY_ENTRY = dt.time(15, 20)
CANDLE_SECONDS = 15 * 60
DEFAULT_POLL_SECONDS = 5 * 60
HISTORY_DAYS = 20
MAX_API_CALLS_PER_SECOND = 8
MINIMUM_CANDLES = 50

ORDER_TYPE = 2  # FYERS: 2 = MARKET
ORDER_SIDE = 1  # FYERS: 1 = BUY
PRODUCT_TYPE = "INTRADAY"

sys.path.insert(0, str(STRATEGIES_DIR / "utils"))
sys.path.insert(0, str(REPO_ROOT / "skills" / "fyers-trading" / "scripts"))
sys.path.insert(0, str(STRATEGIES_DIR / "config"))

_indicator_spec = importlib.util.spec_from_file_location(
    "equity_cross_common_indicators",
    STRATEGIES_DIR / "utils" / "common_indicators.py",
)
assert _indicator_spec and _indicator_spec.loader
_indicator_module = importlib.util.module_from_spec(_indicator_spec)
_indicator_spec.loader.exec_module(_indicator_module)
ema = _indicator_module.ema

from fyers_client import AmbiguousOrderError, FyersAuthError  # noqa: E402
from order_config import (  # noqa: E402
    ConfigGatedFyersClient,
    DailyStockLimitError,
    OrderPlacementDisabledError,
    get_entry_qty,
)


@dataclass(frozen=True)
class Candle:
    """One FYERS OHLCV candle; epoch is the candle start time in seconds."""

    epoch: int
    open: float
    high: float
    low: float
    close: float
    volume: float


SignalEvaluator = Callable[[list[Candle]], tuple[bool, dict]]


class ApiRateLimiter:
    """Space FYERS history calls below the configured requests-per-second cap."""

    def __init__(self, calls_per_second: int = MAX_API_CALLS_PER_SECOND):
        if calls_per_second <= 0:
            raise ValueError("calls_per_second must be positive")
        self.interval = 1.0 / calls_per_second
        self.last_call = 0.0

    def call(self, function, *args, **kwargs):
        elapsed = time.monotonic() - self.last_call
        if elapsed < self.interval:
            time.sleep(self.interval - elapsed)
        self.last_call = time.monotonic()
        return function(*args, **kwargs)


@contextmanager
def exclusive_file_lock(path: pathlib.Path):
    """Serialize cross-process workbook updates on Windows and POSIX."""
    lock_path = path.with_suffix(path.suffix + ".lock")
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    with lock_path.open("a+b") as lock_file:
        lock_file.seek(0, os.SEEK_END)
        if lock_file.tell() == 0:
            lock_file.write(b"\0")
            lock_file.flush()
        lock_file.seek(0)

        if os.name == "nt":
            import msvcrt

            while True:
                try:
                    msvcrt.locking(lock_file.fileno(), msvcrt.LK_NBLCK, 1)
                    break
                except OSError:
                    time.sleep(0.1)
        else:
            import fcntl

            fcntl.flock(lock_file.fileno(), fcntl.LOCK_EX)

        try:
            yield
        finally:
            lock_file.seek(0)
            if os.name == "nt":
                msvcrt.locking(lock_file.fileno(), msvcrt.LK_UNLCK, 1)
            else:
                fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)


def order_tag_for_signal(signal_id: str) -> str:
    """Build a deterministic alphanumeric FYERS orderTag (maximum 30 chars)."""
    digest = hashlib.sha256(signal_id.encode("utf-8")).hexdigest()[:20]
    return f"e15x{digest}"


def signal_lock_path(
    signal_id: str,
    directory: pathlib.Path | None = None,
) -> pathlib.Path:
    """Return a per-signal lock path used to serialize crash recovery."""
    digest = hashlib.sha256(signal_id.encode("utf-8")).hexdigest()[:24]
    base = directory or (STRATEGIES_DIR / "databases")
    return base / f"equity_signal_{digest}"


class SignalStore:
    """Durable state machine and exactly-once workbook identity ledger."""

    def __init__(self, path: pathlib.Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.execute("""
                CREATE TABLE IF NOT EXISTS matched_signals (
                    signal_id TEXT PRIMARY KEY,
                    strategy_name TEXT NOT NULL,
                    symbol TEXT NOT NULL,
                    candle_epoch INTEGER NOT NULL,
                    candle_time_ist TEXT NOT NULL,
                    matched_at_ist TEXT NOT NULL,
                    details_json TEXT NOT NULL,
                    mode TEXT NOT NULL,
                    quantity INTEGER NOT NULL,
                    order_tag TEXT NOT NULL DEFAULT '',
                    order_status TEXT NOT NULL DEFAULT 'MATCHED',
                    order_id TEXT NOT NULL DEFAULT '',
                    broker_message TEXT NOT NULL DEFAULT '',
                    match_printed INTEGER NOT NULL DEFAULT 0,
                    attempted_at_ist TEXT NOT NULL DEFAULT '',
                    updated_at_ist TEXT NOT NULL DEFAULT '',
                    excel_logged INTEGER NOT NULL DEFAULT 0,
                    excel_needs_update INTEGER NOT NULL DEFAULT 1,
                    excel_revision INTEGER NOT NULL DEFAULT 0,
                    UNIQUE (strategy_name, symbol, candle_epoch)
                )
            """)
            existing_columns = {
                row[1] for row in connection.execute("PRAGMA table_info(matched_signals)")
            }
            legacy_pending_state = "order_tag" not in existing_columns
            migrations = {
                "order_tag": "TEXT NOT NULL DEFAULT ''",
                "match_printed": "INTEGER NOT NULL DEFAULT 0",
                "attempted_at_ist": "TEXT NOT NULL DEFAULT ''",
                "updated_at_ist": "TEXT NOT NULL DEFAULT ''",
                "excel_needs_update": "INTEGER NOT NULL DEFAULT 1",
                "excel_revision": "INTEGER NOT NULL DEFAULT 0",
            }
            for column, definition in migrations.items():
                if column not in existing_columns:
                    connection.execute(
                        f"ALTER TABLE matched_signals ADD COLUMN {column} {definition}"
                    )
            if legacy_pending_state:
                # The first version used PENDING for both pre-send and post-send
                # states and a shared orderTag. Never auto-send these rows.
                connection.execute(
                    """
                    UPDATE matched_signals
                    SET order_status = 'UNKNOWN_LEGACY',
                        order_tag = 'legacy_static_tag',
                        broker_message = 'legacy pending state; manual broker review required'
                    WHERE order_status = 'PENDING'
                    """
                )
            blank_tag_rows = connection.execute(
                "SELECT signal_id FROM matched_signals WHERE order_tag = ''"
            ).fetchall()
            for (signal_id,) in blank_tag_rows:
                connection.execute(
                    "UPDATE matched_signals SET order_tag = ? WHERE signal_id = ?",
                    (order_tag_for_signal(signal_id), signal_id),
                )
            connection.commit()

    def claim(
        self,
        symbol: str,
        candle_epoch: int,
        candle_time_ist: str,
        matched_at_ist: str,
        details: dict,
        mode: str,
        quantity: int,
        strategy_name: str = STRATEGY_NAME,
    ) -> str | None:
        """Atomically return a signal ID only for the first matching process."""
        signal_id = f"{strategy_name}:{symbol}:{candle_epoch}"
        with closing(sqlite3.connect(self.path)) as connection, connection:
            cursor = connection.execute(
                """
                INSERT OR IGNORE INTO matched_signals
                    (signal_id, strategy_name, symbol, candle_epoch,
                     candle_time_ist, matched_at_ist, details_json, mode,
                     quantity, order_tag, updated_at_ist)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    signal_id,
                    strategy_name,
                    symbol,
                    candle_epoch,
                    candle_time_ist,
                    matched_at_ist,
                    json.dumps(details, sort_keys=True),
                    mode,
                    quantity,
                    order_tag_for_signal(signal_id),
                    matched_at_ist,
                ),
            )
            return signal_id if cursor.rowcount == 1 else None

    def finish_order(
        self,
        signal_id: str,
        order_status: str,
        order_id: str,
        broker_message: str,
    ) -> None:
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.execute(
                """
                UPDATE matched_signals
                SET order_status = ?, order_id = ?, broker_message = ?,
                    updated_at_ist = ?, excel_needs_update = 1,
                    excel_revision = excel_revision + 1
                WHERE signal_id = ?
                """,
                (
                    order_status,
                    order_id,
                    broker_message,
                    now_ist().isoformat(timespec="seconds"),
                    signal_id,
                ),
            )
            connection.commit()

    def mark_match_printed(self, signal_id: str) -> None:
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.execute(
                "UPDATE matched_signals SET match_printed = 1 WHERE signal_id = ?",
                (signal_id,),
            )
            connection.commit()

    def mark_submitting(self, signal_id: str) -> bool:
        """Own the one pre-send transition; only the winner may transmit."""
        timestamp = now_ist().isoformat(timespec="seconds")
        with closing(sqlite3.connect(self.path)) as connection, connection:
            cursor = connection.execute(
                """
                UPDATE matched_signals
                SET order_status = 'SUBMITTING', attempted_at_ist = ?,
                    updated_at_ist = ?, excel_needs_update = 1,
                    excel_revision = excel_revision + 1
                WHERE signal_id = ? AND order_status IN ('MATCHED', 'PENDING')
                """,
                (timestamp, timestamp, signal_id),
            )
            return cursor.rowcount == 1

    def unfinished_rows(self) -> list[sqlite3.Row]:
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.row_factory = sqlite3.Row
            return connection.execute(
                """
                SELECT * FROM matched_signals
                WHERE order_status IN (
                    'MATCHED', 'PENDING', 'SUBMITTING', 'UNKNOWN', 'PENDING_ACK'
                )
                ORDER BY candle_epoch, symbol
                """
            ).fetchall()

    def pending_excel_rows(self) -> list[sqlite3.Row]:
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.row_factory = sqlite3.Row
            return connection.execute(
                """
                SELECT * FROM matched_signals
                WHERE excel_needs_update = 1
                ORDER BY candle_epoch, symbol
                """
            ).fetchall()

    def get_row(self, signal_id: str) -> sqlite3.Row:
        with closing(sqlite3.connect(self.path)) as connection, connection:
            connection.row_factory = sqlite3.Row
            row = connection.execute(
                "SELECT * FROM matched_signals WHERE signal_id = ?",
                (signal_id,),
            ).fetchone()
        if row is None:
            raise RuntimeError(f"signal ledger row not found: {signal_id}")
        return row

    def mark_excel_synced(self, signal_id: str, revision: int) -> bool:
        """Acknowledge only the exact SQLite snapshot written to Excel."""
        with closing(sqlite3.connect(self.path)) as connection, connection:
            cursor = connection.execute(
                """
                UPDATE matched_signals
                SET excel_logged = 1, excel_needs_update = 0
                WHERE signal_id = ? AND excel_revision = ?
                """,
                (signal_id, revision),
            )
            return cursor.rowcount == 1


EXCEL_HEADERS = [
    "SignalID",
    "MatchedAtIST",
    "SignalCandleStartIST",
    "SignalCandleEndIST",
    "Symbol",
    "Timeframe",
    "Open",
    "High",
    "Low",
    "Close",
    "PreviousClose",
    "EMA15",
    "EMA10",
    "EMA20",
    "EMA50",
    "Mode",
    "Quantity",
    "OrderType",
    "ProductType",
    "OrderStatus",
    "OrderID",
    "OrderTag",
    "BrokerMessage",
    "Details",
]
LEGACY_EXCEL_HEADERS = [
    header for header in EXCEL_HEADERS if header != "OrderTag"
]
EXCEL_COLUMN_WIDTHS = {
    "A": 65,
    "B": 27,
    "C": 27,
    "D": 27,
    "E": 24,
    "F": 12,
    "G": 13,
    "H": 13,
    "I": 13,
    "J": 13,
    "K": 16,
    "L": 14,
    "M": 14,
    "N": 14,
    "O": 14,
    "P": 12,
    "Q": 12,
    "R": 12,
    "S": 15,
    "T": 16,
    "U": 25,
    "V": 28,
    "W": 45,
    "X": 90,
}


def now_ist() -> dt.datetime:
    return dt.datetime.now(MARKET_TIMEZONE)


def log_message(message: str, error: bool = False) -> None:
    """Append one timestamped event to the console and strategy log."""
    entry = f"{now_ist().isoformat(timespec='seconds')} {message}"
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOG_PATH.open("a", encoding="utf-8") as log_file:
        log_file.write(entry + "\n")
    print(entry, file=sys.stderr if error else sys.stdout)


def read_symbols(path: pathlib.Path = STOCKS_PATH) -> list[str]:
    """Read one FYERS equity symbol per non-comment line."""
    symbols = [
        line.strip()
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip() and not line.lstrip().startswith("#")
    ]
    return list(dict.fromkeys(symbols))


def completed_candles_from_response(
    response: dict,
    current_time: dt.datetime | None = None,
) -> list[Candle]:
    """Parse, deduplicate, sort, and retain only completed 15-minute candles."""
    current = current_time or now_ist()
    if current.tzinfo is None:
        current = current.replace(tzinfo=MARKET_TIMEZONE)
    current_epoch = int(current.timestamp())
    by_epoch: dict[int, Candle] = {}

    for row in response.get("candles") or []:
        if not isinstance(row, (list, tuple)) or len(row) < 6:
            continue
        try:
            candle = Candle(int(row[0]), *map(float, row[1:6]))
        except (TypeError, ValueError):
            continue
        if candle.epoch + CANDLE_SECONDS <= current_epoch:
            by_epoch[candle.epoch] = candle

    return [by_epoch[epoch] for epoch in sorted(by_epoch)]


def fetch_completed_candles(
    client,
    symbol: str,
    limiter: ApiRateLimiter,
    current_time: dt.datetime | None = None,
) -> list[Candle]:
    """Fetch history and return only candles whose full 15 minutes have elapsed."""
    current = current_time or now_ist()
    if current.tzinfo is None:
        current = current.replace(tzinfo=MARKET_TIMEZONE)
    start_date = current.date() - dt.timedelta(days=HISTORY_DAYS)
    response = limiter.call(
        client.history,
        symbol,
        "15",
        start_date.isoformat(),
        current.date().isoformat(),
    )
    if response.get("s") != "ok":
        raise RuntimeError(f"history fetch failed for {symbol}: {response}")
    return completed_candles_from_response(response, current)


def epoch_is_current(candle_epoch: int, current: dt.datetime) -> bool:
    """Return whether a completed candle is the latest expected 15-minute bar."""
    if current.tzinfo is None:
        current = current.replace(tzinfo=MARKET_TIMEZONE)
    candle_end = candle_epoch + CANDLE_SECONDS
    return 0 <= current.timestamp() - candle_end < CANDLE_SECONDS


def candle_is_current(candle: Candle, current: dt.datetime) -> bool:
    """Reject delayed history so an old crossover cannot trigger a late entry."""
    return epoch_is_current(candle.epoch, current)


def evaluate_fresh_crossover(candles: list[Candle]) -> tuple[bool, dict]:
    """Evaluate all entry rules on the latest two completed candles."""
    if len(candles) < MINIMUM_CANDLES + 1:
        return False, {"reason": f"need at least {MINIMUM_CANDLES + 1} completed candles"}

    closes = [candle.close for candle in candles]
    indicator_series = {period: ema(closes, period) for period in (10, 15, 20, 50)}
    previous = candles[-2]
    current = candles[-1]

    required_values = [
        series[-2]
        for series in indicator_series.values()
    ] + [series[-1] for series in indicator_series.values()]
    if any(value is None for value in required_values):
        return False, {"reason": "insufficient EMA history"}

    ema10_values = indicator_series[10]
    ema15_values = indicator_series[15]
    ema20_values = indicator_series[20]
    ema50_values = indicator_series[50]

    previous_ema10 = ema10_values[-2]
    current_ema10 = ema10_values[-1]
    previous_ema15 = ema15_values[-2]
    current_ema15 = ema15_values[-1]
    previous_ema20 = ema20_values[-2]
    current_ema20 = ema20_values[-1]
    previous_ema50 = ema50_values[-2]
    current_ema50 = ema50_values[-1]

    # Equality counts as the pre-cross side. An exact match therefore needs
    # previous <= fast/slow and current > fast/slow.
    conditions = {
        "close_crossed_above_ema15": (
            previous.close <= previous_ema15 and current.close > current_ema15
        ),
        "close_above_ema10": current.close > current_ema10,
        "ema10_crossed_above_ema20": (
            previous_ema10 <= previous_ema20 and current_ema10 > current_ema20
        ),
        "ema20_crossed_above_ema50": (
            previous_ema20 <= previous_ema50 and current_ema20 > current_ema50
        ),
        "not_previously_ema10_ema20_ema50_bullish_stack": not (
            previous_ema10 > previous_ema20 > previous_ema50
        ),
    }
    details = {
        "candle_epoch": current.epoch,
        "open": current.open,
        "high": current.high,
        "low": current.low,
        "close": current.close,
        "previous_close": previous.close,
        "ema15": current_ema15,
        "ema10": current_ema10,
        "ema20": current_ema20,
        "ema50": current_ema50,
        "previous_ema15": previous_ema15,
        "previous_ema10": previous_ema10,
        "previous_ema20": previous_ema20,
        "previous_ema50": previous_ema50,
        "conditions": conditions,
    }
    triggered = all(conditions.values())
    if not triggered:
        details["failed_conditions"] = [
            name for name, passed in conditions.items() if not passed
        ]
    return triggered, details


def evaluate_ema10_pullback_cross(candles: list[Candle]) -> tuple[bool, dict]:
    """Evaluate the independent close-upcross/aligned-EMA strategy."""
    if len(candles) < MINIMUM_CANDLES + 1:
        return False, {"reason": f"need at least {MINIMUM_CANDLES + 1} completed candles"}

    closes = [candle.close for candle in candles]
    indicator_series = {period: ema(closes, period) for period in (10, 15, 20, 50)}
    previous = candles[-2]
    current = candles[-1]
    required_values = [
        series[-2]
        for series in indicator_series.values()
    ] + [series[-1] for series in indicator_series.values()]
    if any(value is None for value in required_values):
        return False, {"reason": "insufficient EMA history"}

    ema10_values = indicator_series[10]
    ema15_values = indicator_series[15]
    ema20_values = indicator_series[20]
    ema50_values = indicator_series[50]
    previous_ema10 = ema10_values[-2]
    current_ema10 = ema10_values[-1]
    previous_ema15 = ema15_values[-2]
    current_ema15 = ema15_values[-1]
    previous_ema20 = ema20_values[-2]
    current_ema20 = ema20_values[-1]
    previous_ema50 = ema50_values[-2]
    current_ema50 = ema50_values[-1]

    conditions = {
        "previous_close_below_ema10": previous.close < previous_ema10,
        "current_close_above_ema10": current.close > current_ema10,
        "not_crossed_from_above_to_below_ema10": not (
            previous.close > previous_ema10 and current.close < current_ema10
        ),
        "ema10_above_ema20": current_ema10 > current_ema20,
        "ema20_above_ema50": current_ema20 > current_ema50,
        "ema10_ema20_ema50_bullish_stack": (
            current_ema10 > current_ema20 > current_ema50
        ),
    }
    details = {
        "candle_epoch": current.epoch,
        "open": current.open,
        "high": current.high,
        "low": current.low,
        "close": current.close,
        "previous_close": previous.close,
        "ema15": current_ema15,
        "ema10": current_ema10,
        "ema20": current_ema20,
        "ema50": current_ema50,
        "previous_ema15": previous_ema15,
        "previous_ema10": previous_ema10,
        "previous_ema20": previous_ema20,
        "previous_ema50": previous_ema50,
        "conditions": conditions,
    }
    triggered = all(conditions.values())
    if not triggered:
        details["failed_conditions"] = [
            name for name, passed in conditions.items() if not passed
        ]
    return triggered, details


def order_result(response: dict) -> tuple[str, str, str]:
    """Normalize a FYERS order response for logs and Excel."""
    broker_status = str(response.get("s", "")).strip().lower()
    order_id = str(response.get("id") or response.get("id_fyers") or "")
    message = str(response.get("message", ""))
    code = response.get("code")

    if broker_status == "dry_run":
        status = "DRY_RUN"
        order_id = order_id or "DRY_RUN"
    elif broker_status == "ok" and str(code) == "201":
        status = "PENDING_ACK"
        order_id = order_id or "PENDING_ACK"
    elif broker_status == "ok":
        status = "PLACED"
        order_id = order_id or "NOT_RETURNED"
    elif broker_status == "error" or (isinstance(code, int) and code < 0):
        status = "REJECTED"
        order_id = order_id or "NOT_RETURNED"
    else:
        status = broker_status.upper() or "UNKNOWN"
        order_id = order_id or "NOT_RETURNED"
    return status, order_id, message


def print_match(
    signal_id: str,
    symbol: str,
    matched_at: dt.datetime,
    details: dict,
) -> None:
    """Print full stock/candle/EMA details for one signal ID."""
    log_message(
        "MATCH | "
        f"signal_id={signal_id} | strategy={details.get('strategy_name', 'unknown')} | "
        f"symbol={symbol} | timeframe=15m | "
        f"matched_at={matched_at.isoformat(timespec='seconds')} | "
        f"candle_start={dt.datetime.fromtimestamp(details['candle_epoch'], MARKET_TIMEZONE).isoformat(timespec='seconds')} | "
        f"O={details['open']:.2f} H={details['high']:.2f} "
        f"L={details['low']:.2f} C={details['close']:.2f} | "
        f"EMA15={details['ema15']:.2f} EMA10={details['ema10']:.2f} "
        f"EMA20={details['ema20']:.2f} EMA50={details['ema50']:.2f} | "
        f"conditions={json.dumps(details['conditions'], sort_keys=True)}"
    )


def _create_excel_workbook(path: pathlib.Path) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Fresh_Crossovers"
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for column, header in enumerate(EXCEL_HEADERS, 1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for column, width in EXCEL_COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width
    worksheet.freeze_panes = "A2"
    return workbook


def _excel_values(row: sqlite3.Row) -> list:
    details = json.loads(row["details_json"])
    return [
        row["signal_id"],
        row["matched_at_ist"],
        row["candle_time_ist"],
        dt.datetime.fromtimestamp(
            row["candle_epoch"] + CANDLE_SECONDS, MARKET_TIMEZONE
        ).isoformat(timespec="seconds"),
        row["symbol"],
        "15m",
        details["open"],
        details["high"],
        details["low"],
        details["close"],
        details["previous_close"],
        details["ema15"],
        details["ema10"],
        details["ema20"],
        details["ema50"],
        row["mode"],
        row["quantity"],
        "MARKET",
        PRODUCT_TYPE,
        row["order_status"],
        row["order_id"],
        row["order_tag"],
        row["broker_message"],
        row["details_json"],
    ]


def upsert_signal_to_excel(path: pathlib.Path, row: sqlite3.Row) -> bool:
    """Create/update one SignalID row under a cross-process file lock."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with exclusive_file_lock(path):
        workbook = None
        temporary_path: pathlib.Path | None = None
        try:
            if path.exists():
                workbook = load_workbook(path)
                worksheet = workbook.active
                headers = [
                    worksheet.cell(1, column).value
                    for column in range(1, worksheet.max_column + 1)
                ]
                if headers == LEGACY_EXCEL_HEADERS:
                    worksheet.insert_cols(22, 1)
                    worksheet.cell(1, 22, "OrderTag")
                    for row_number in range(2, worksheet.max_row + 1):
                        signal_id = worksheet.cell(row_number, 1).value
                        worksheet.cell(
                            row_number,
                            22,
                            order_tag_for_signal(str(signal_id)) if signal_id else "",
                        )
                    headers = [
                        worksheet.cell(1, column).value
                        for column in range(1, len(EXCEL_HEADERS) + 1)
                    ]
                if headers != EXCEL_HEADERS:
                    raise RuntimeError(f"unexpected Excel header in {path}")
            else:
                workbook = _create_excel_workbook(path)
                worksheet = workbook.active

            output_row = None
            for row_number in range(2, worksheet.max_row + 1):
                if worksheet.cell(row_number, 1).value == row["signal_id"]:
                    output_row = row_number
                    break
            if output_row is None:
                output_row = worksheet.max_row + 1

            for column, value in enumerate(_excel_values(row), 1):
                cell = worksheet.cell(row=output_row, column=column, value=value)
                cell.alignment = Alignment(
                    vertical="top", wrap_text=column == len(EXCEL_HEADERS)
                )
            worksheet.auto_filter.ref = (
                f"A1:{worksheet.cell(worksheet.max_row, len(EXCEL_HEADERS)).coordinate}"
            )

            descriptor, temporary_name = tempfile.mkstemp(
                prefix=f".{path.stem}.",
                suffix=".tmp.xlsx",
                dir=path.parent,
            )
            os.close(descriptor)
            temporary_path = pathlib.Path(temporary_name)
            workbook.save(temporary_path)
            workbook.close()
            workbook = None
            os.replace(temporary_path, path)
            temporary_path = None
            return True
        finally:
            if workbook is not None:
                workbook.close()
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)


def append_signal_to_excel(path: pathlib.Path, row: sqlite3.Row) -> bool:
    """Backward-compatible name for the workbook upsert operation."""
    return upsert_signal_to_excel(path, row)


def flush_pending_excel(store: SignalStore, path: pathlib.Path) -> bool:
    """Upsert pending rows; keep scanning if Excel is open or unavailable."""
    for row in store.pending_excel_rows():
        try:
            upsert_signal_to_excel(path, row)
            if not store.mark_excel_synced(
                row["signal_id"], row["excel_revision"]
            ):
                # SQLite advanced while this workbook snapshot was being saved.
                # Leave excel_needs_update=1 so the newer row is written next.
                continue
        except Exception as error:
            log_message(
                f"EXCEL DEFERRED | {row['signal_id']} | {error}",
                error=True,
            )
            return False
    return True


def _walk_order_items(value):
    """Yield order-like dictionaries from common FYERS orderbook envelopes."""
    if isinstance(value, list):
        for item in value:
            yield from _walk_order_items(item)
    elif isinstance(value, dict):
        order_tag = value.get("orderTag") or value.get("order_tag")
        if order_tag:
            yield value
        for key in ("d", "orderBook", "orderbook", "orders"):
            if key in value:
                yield from _walk_order_items(value[key])


def find_order_by_tag(response: dict, order_tag: str) -> dict | None:
    """Find the broker order associated with this signal's deterministic tag."""
    if response.get("s") != "ok":
        return None
    return next(
        (
            item
            for item in _walk_order_items(response)
            if str(item.get("orderTag") or item.get("order_tag")) == order_tag
        ),
        None,
    )


def orderbook_order_result(item: dict) -> tuple[str, str, str]:
    order_id = str(
        item.get("id")
        or item.get("orderId")
        or item.get("orderNumber")
        or "FOUND_WITHOUT_ID"
    )
    broker_status = str(
        item.get("orderStatus") or item.get("status") or item.get("ordStatus") or "UNKNOWN"
    )
    return "RECONCILED", order_id, f"orderbook tag match; broker_status={broker_status}"


def reconcile_order(client, store: SignalStore, row: sqlite3.Row) -> bool:
    """Resolve an ambiguous/pending send by tag; never issue another POST."""
    try:
        response = client.orderbook()
    except FyersAuthError:
        raise
    except Exception as error:
        log_message(
            f"RECONCILE FAILED | {row['signal_id']} | {error}",
            error=True,
        )
        return False

    item = find_order_by_tag(response, row["order_tag"])
    if item is None:
        return False
    status, order_id, message = orderbook_order_result(item)
    store.finish_order(row["signal_id"], status, order_id, message)
    log_message(
        f"RECONCILED | {row['symbol']} | order_id={order_id} | {message}"
    )
    return True


def place_buy_order(client, row: sqlite3.Row) -> dict:
    """Submit one tagged order with transient-write retries disabled."""
    live = row["mode"] == "LIVE"
    order = {
        "symbol": row["symbol"],
        "qty": row["quantity"],
        "type": ORDER_TYPE,
        "side": ORDER_SIDE,
        "productType": PRODUCT_TYPE,
        "orderTag": row["order_tag"],
    }
    meta = {
        "strategy": row["strategy_name"],
        "signal": "ENTRY",
        "timeframe": "15m",
        "signal_id": row["signal_id"],
        "description": "Fresh EMA15/10/20/50 crossover entry",
    }
    return client.place_order(
        order,
        dry_run=not live,
        meta=meta,
        retry_transient=False,
    )


def log_contains_signal(signal_id: str) -> bool:
    try:
        return f"signal_id={signal_id}" in LOG_PATH.read_text(encoding="utf-8")
    except OSError:
        return False


def ensure_match_printed(store: SignalStore, row: sqlite3.Row) -> sqlite3.Row:
    if not row["match_printed"]:
        if not log_contains_signal(row["signal_id"]):
            # Do not silently mark a signal printed when the durable log/console
            # write failed; recovery will retry this output.
            print_match(
                row["signal_id"],
                row["symbol"],
                dt.datetime.fromisoformat(row["matched_at_ist"]),
                json.loads(row["details_json"]),
            )
        store.mark_match_printed(row["signal_id"])
        row = store.get_row(row["signal_id"])
    return row


def log_terminal_order(row: sqlite3.Row) -> None:
    log_message(
        f"ORDER | {row['symbol']} | mode={row['mode']} | "
        f"status={row['order_status']} | order_id={row['order_id']} | "
        f"order_tag={row['order_tag']} | message={row['broker_message']}"
    )


def _process_claimed_signal_locked(
    client,
    store: SignalStore,
    row: sqlite3.Row,
    allow_live_send: bool,
) -> None:
    signal_id = row["signal_id"]
    live = row["mode"] == "LIVE"
    row = ensure_match_printed(store, row)

    if live and not allow_live_send:
        # A default dry-run invocation may recover/print the event, but must
        # never turn a persisted LIVE intent into a real order.
        return

    current = now_ist()
    signal_day = dt.datetime.fromtimestamp(
        row["candle_epoch"], MARKET_TIMEZONE
    ).date()
    if live and (
        signal_day != current.date()
        or not epoch_is_current(row["candle_epoch"], current)
    ):
        store.finish_order(
            signal_id,
            "SKIPPED_STALE_RECOVERY",
            "NOT_SENT",
            f"signal candle {signal_day} is not the current completed 15-minute bar",
        )
        log_terminal_order(store.get_row(signal_id))
        return

    if live and not live_order_window_is_open(current):
        store.finish_order(
            signal_id,
            "SKIPPED_LATE_ENTRY",
            "NOT_SENT",
            f"live entry window closed at {LAST_INTRADAY_ENTRY.isoformat(timespec='minutes')} IST",
        )
        log_terminal_order(store.get_row(signal_id))
        return

    owns_processing = store.mark_submitting(signal_id)
    if not owns_processing:
        # Another process already owns or completed this transition.
        return
    row = store.get_row(signal_id)

    auth_error: FyersAuthError | None = None
    try:
        response = place_buy_order(client, row)
        status, order_id, message = order_result(response)
        store.finish_order(signal_id, status, order_id, message)
        if status == "PENDING_ACK":
            reconcile_order(client, store, store.get_row(signal_id))
    except AmbiguousOrderError as error:
        store.finish_order(
            signal_id,
            "UNKNOWN",
            "RECONCILE_REQUIRED",
            f"order transmission ambiguous; not retried: {error}",
        )
        reconcile_order(client, store, store.get_row(signal_id))
    except FyersAuthError as error:
        store.finish_order(signal_id, "FAILED_NOT_SENT", "NOT_SENT", str(error))
        auth_error = error
    except (OrderPlacementDisabledError, DailyStockLimitError, ValueError) as error:
        store.finish_order(signal_id, "NOT_SENT", "NOT_SENT", str(error))
    except RuntimeError as error:
        # Non-ambiguous HTTP 4xx/429 errors from the no-retry write are definite.
        store.finish_order(signal_id, "REJECTED_NOT_SENT", "NOT_SENT", str(error))
    except Exception as error:
        store.finish_order(
            signal_id,
            "UNKNOWN",
            "RECONCILE_REQUIRED",
            f"unexpected order error; not retried: {error}",
        )
        reconcile_order(client, store, store.get_row(signal_id))

    log_terminal_order(store.get_row(signal_id))
    if auth_error is not None:
        raise auth_error


def process_claimed_signal(
    client,
    store: SignalStore,
    row: sqlite3.Row,
    allow_live_send: bool = True,
) -> None:
    """Serialize one signal's print/send/reconcile state machine."""
    signal_id = row["signal_id"]
    with exclusive_file_lock(signal_lock_path(signal_id, store.path.parent)):
        current_row = store.get_row(signal_id)
        _process_claimed_signal_locked(
            client,
            store,
            current_row,
            allow_live_send,
        )


def recover_unfinished_signals(
    client,
    store: SignalStore,
    excel_path: pathlib.Path,
    allow_live_send: bool = True,
) -> None:
    """Resume safe pre-send work and reconcile every ambiguous/pending send."""
    for snapshot in store.unfinished_rows():
        with exclusive_file_lock(
            signal_lock_path(snapshot["signal_id"], store.path.parent)
        ):
            row = store.get_row(snapshot["signal_id"])
            row = ensure_match_printed(store, row)
            status = row["order_status"]
            if row["mode"] == "DRY_RUN" and status == "SUBMITTING":
                store.finish_order(
                    row["signal_id"],
                    "DRY_RUN_INCOMPLETE",
                    "NOT_SENT",
                    "dry-run process stopped before completion; no live order was sent",
                )
            elif status in ("MATCHED", "PENDING"):
                _process_claimed_signal_locked(
                    client,
                    store,
                    row,
                    allow_live_send,
                )
            else:
                reconciled = reconcile_order(client, store, row)
                if not reconciled and status == "SUBMITTING":
                    store.finish_order(
                        row["signal_id"],
                        "UNKNOWN",
                        "RECONCILE_REQUIRED",
                        "submission ownership was persisted but no tagged order is visible; automatic resend is disabled",
                    )
    flush_pending_excel(store, excel_path)


def handle_match(
    client,
    store: SignalStore,
    symbol: str,
    candles: list[Candle],
    live: bool,
    excel_path: pathlib.Path = EXCEL_PATH,
    strategy_name: str = STRATEGY_NAME,
    evaluator: SignalEvaluator = evaluate_fresh_crossover,
) -> bool:
    """Claim, process, and Excel-upsert one strategy match at most once."""
    triggered, details = evaluator(candles)
    if not triggered:
        return False

    details = dict(details)
    details["strategy_name"] = strategy_name
    matched_at = now_ist()
    candle_start = dt.datetime.fromtimestamp(
        details["candle_epoch"], MARKET_TIMEZONE
    )
    signal_id = store.claim(
        symbol,
        details["candle_epoch"],
        candle_start.isoformat(timespec="seconds"),
        matched_at.isoformat(timespec="seconds"),
        details,
        "LIVE" if live else "DRY_RUN",
        get_entry_qty(SCRIPT_NAME),
        strategy_name=strategy_name,
    )
    if signal_id is None:
        return False

    process_claimed_signal(
        client,
        store,
        store.get_row(signal_id),
        allow_live_send=live,
    )
    flush_pending_excel(store, excel_path)
    return True


def market_is_open(current: dt.datetime | None = None) -> bool:
    current = current or now_ist()
    return (
        current.weekday() < 5
        and MARKET_OPEN <= current.time() <= MARKET_CLOSE
    )


def live_order_window_is_open(current: dt.datetime | None = None) -> bool:
    """Keep new MIS orders before the broker's late-session square-off window."""
    current = current or now_ist()
    return market_is_open(current) and current.time() <= LAST_INTRADAY_ENTRY


def seconds_until_next_candle_close(current: dt.datetime | None = None) -> float:
    """Return the delay to the next 15-minute boundary, never sleeping past it."""
    current = current or now_ist()
    if current.tzinfo is None:
        current = current.replace(tzinfo=MARKET_TIMEZONE)
    next_boundary = ((current.minute // 15) + 1) * 15
    if next_boundary == 60:
        target = current.replace(
            minute=0, second=0, microsecond=0
        ) + dt.timedelta(hours=1)
    else:
        target = current.replace(
            minute=next_boundary, second=0, microsecond=0
        )
    if target <= current:
        target += dt.timedelta(minutes=15)
    return max(1.0, (target - current).total_seconds())


STRATEGY_DEFINITIONS = (
    (STRATEGY_NAME, evaluate_fresh_crossover),
    (STRATEGY_NAME_PULLBACK, evaluate_ema10_pullback_cross),
)


def run_cycle(
    client,
    limiter: ApiRateLimiter,
    store: SignalStore,
    live: bool,
    symbols: list[str],
    excel_path: pathlib.Path = EXCEL_PATH,
) -> None:
    """Recover durable work, then fetch and evaluate every symbol once."""
    recover_unfinished_signals(
        client,
        store,
        excel_path,
        allow_live_send=live,
    )
    for symbol in symbols:
        current = now_ist()
        try:
            candles = fetch_completed_candles(client, symbol, limiter, current)
            if not candles:
                continue
            latest = candles[-1]
            latest_time = dt.datetime.fromtimestamp(latest.epoch, MARKET_TIMEZONE)
            if latest_time.date() != current.date() or not candle_is_current(
                latest, current
            ):
                # Never repeat a prior-session signal or act on delayed history.
                continue
            for strategy_name, evaluator in STRATEGY_DEFINITIONS:
                handle_match(
                    client,
                    store,
                    symbol,
                    candles,
                    live,
                    excel_path,
                    strategy_name=strategy_name,
                    evaluator=evaluator,
                )
        except FyersAuthError:
            raise
        except (OSError, RuntimeError, ValueError, sqlite3.Error) as error:
            log_message(f"SCAN ERROR | {symbol} | {error}", error=True)
    flush_pending_excel(store, excel_path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--live",
        action="store_true",
        help="Send real orders. Also requires place_order=YES in the equity config.",
    )
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="Print and build orders without sending them (default).",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="Run one scan cycle and exit (useful outside market hours for dry-run).",
    )
    parser.add_argument(
        "--poll-seconds",
        type=int,
        default=DEFAULT_POLL_SECONDS,
        help=f"Seconds between scans (default: {DEFAULT_POLL_SECONDS}).",
    )
    args = parser.parse_args(argv)
    if args.poll_seconds < 1:
        parser.error("--poll-seconds must be positive")
    return args


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    live = bool(args.live)
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    log_message(
        f"STARTED | strategies={STRATEGY_NAME},{STRATEGY_NAME_PULLBACK} | "
        f"mode={'LIVE' if live else 'DRY_RUN'} | "
        f"market={MARKET_OPEN}-{MARKET_CLOSE} IST | poll={args.poll_seconds}s"
    )

    try:
        client = ConfigGatedFyersClient(
            strategy_name=STRATEGY_NAME,
            script_name=SCRIPT_NAME,
        )
        limiter = ApiRateLimiter()
        store = SignalStore(DATABASE_PATH)
        # Reconcile prior sends even when today's entry window is closed.
        recover_unfinished_signals(
            client,
            store,
            EXCEL_PATH,
            allow_live_send=live,
        )

        if live and not live_order_window_is_open():
            log_message(
                "STOPPED | live entries are allowed only from 09:15 through 15:20 IST",
                error=True,
            )
            return 0

        symbols = read_symbols()
        if not symbols:
            raise ValueError(f"no symbols found in {STOCKS_PATH}")
        log_message(f"Universe | {len(symbols)} NIFTY 50 equity symbols")

        if not args.once and not market_is_open():
            log_message("STOPPED | market is closed")
            return 0

        while args.once or market_is_open():
            run_cycle(client, limiter, store, live, symbols)
            if args.once:
                break
            time.sleep(
                min(args.poll_seconds, seconds_until_next_candle_close())
            )

        log_message("STOPPED | scan complete")
        return 0
    except FyersAuthError as error:
        log_message(f"AUTHENTICATION ERROR | {error}", error=True)
        return 1
    except (OSError, RuntimeError, ValueError, sqlite3.Error) as error:
        log_message(f"ERROR | {error}", error=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
